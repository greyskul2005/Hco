"""Convierte markdown estructurado a DOCX y un PDF ejecutivo."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib.colors import HexColor, white
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    ListFlowable,
    ListItem,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
BLUE = RGBColor(0x1D, 0x4E, 0xD8)
TEAL = RGBColor(0x0E, 0x74, 0x90)
INK = RGBColor(0x0F, 0x17, 0x2A)
MUTED = RGBColor(0x47, 0x55, 0x69)


def _shade(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def _set_run_font(run, name="Calibri", size=11, bold=False, color=INK):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def new_doc(title: str, subtitle: str) -> Document:
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(2.0)
        s.bottom_margin = Cm(2.0)
        s.left_margin = Cm(2.2)
        s.right_margin = Cm(2.2)
        header = s.header
        hp = header.paragraphs[0]
        hp.text = f"HC Control v2.0  ·  {title}  ·  CONFIDENCIAL INTERNO  ·  datos sintéticos"
        hp.runs[0].font.size = Pt(8)
        hp.runs[0].font.color.rgb = MUTED
        footer = s.footer
        fp = footer.paragraphs[0]
        fp.text = "Paquete de implantación · no contiene personas reales · 514 + 13 − 9 + 9 = 527"
        fp.runs[0].font.size = Pt(8)
        fp.runs[0].font.color.rgb = MUTED
    # cover
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("HC CONTROL")
    _set_run_font(r, size=14, bold=True, color=BLUE)
    p = doc.add_paragraph()
    r = p.add_run(title)
    _set_run_font(r, size=28, bold=True, color=NAVY)
    p = doc.add_paragraph()
    r = p.add_run(subtitle)
    _set_run_font(r, size=12, color=TEAL)
    meta = [
        "Versión 2.0  ·  20 de septiembre de 2026",
        "Clasificación: uso interno de implantación",
        "Datos de personas: exclusivamente sintéticos de demostración",
        "Cifras confirmadas: HC base 514 · incrementos +13 · CEV −9/+9 · HC final 527",
        "Power Apps: fuera de alcance. MVP operable con Microsoft Lists, Excel y Teams.",
    ]
    for line in meta:
        p = doc.add_paragraph()
        r = p.add_run(line)
        _set_run_font(r, size=11, color=INK)
    doc.add_paragraph()
    return doc


def add_h(doc: Document, text: str, level: int):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = NAVY
        run.font.name = "Calibri"


def add_p(doc: Document, text: str, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    _set_run_font(r, size=11, color=INK)
    r.italic = italic
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    return p


def add_callout(doc: Document, kind: str, text: str):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    colors = {
        "Confirmado": "D1FAE5",
        "Pendiente": "FEF3C7",
        "Supuesto de diseño": "E0F2FE",
        "Recomendación": "EDE9FE",
        "Incidencia": "FEE2E2",
    }
    _shade(cell, colors.get(kind, "F1F5F9"))
    p = cell.paragraphs[0]
    r = p.add_run(f"{kind.upper()}. ")
    _set_run_font(r, size=10, bold=True, color=NAVY)
    r2 = p.add_run(text)
    _set_run_font(r2, size=10, color=INK)
    doc.add_paragraph()


def add_table(doc: Document, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        r = p.add_run(str(h))
        _set_run_font(r, size=9, bold=True, color=RGBColor(255, 255, 255))
        _shade(hdr[i], "0B1F3A")
    for row in rows:
        cells = table.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            r = p.add_run(str(v))
            _set_run_font(r, size=9, color=INK)
            if len(table.rows) % 2 == 0:
                _shade(cells[i], "F8FAFC")
    doc.add_paragraph()


def add_bullets(doc: Document, items: list[str]):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        p.clear()
        r = p.add_run(it)
        _set_run_font(r, size=11)


def markdown_to_docx(md: str, dest: Path, title: str, subtitle: str):
    doc = new_doc(title, subtitle)
    lines = md.splitlines()
    i = 0
    table_buf = []
    in_table = False
    para_buf: list[str] = []

    def flush_p():
        nonlocal para_buf
        if para_buf:
            add_p(doc, " ".join(para_buf).strip())
            para_buf = []

    def flush_table():
        nonlocal table_buf, in_table
        if not table_buf:
            in_table = False
            return
        parsed = []
        for raw in table_buf:
            cols = [c.strip() for c in raw.strip().strip("|").split("|")]
            parsed.append(cols)
        if len(parsed) >= 2:
            headers = parsed[0]
            rows = [r for r in parsed[2:] if any(x.strip() for x in r)]
            add_table(doc, headers, rows)
        table_buf = []
        in_table = False

    while i < len(lines):
        line = lines[i]
        if line.startswith("|"):
            flush_p()
            in_table = True
            table_buf.append(line)
            i += 1
            continue
        if in_table:
            flush_table()
        if line.startswith("# "):
            flush_p()
            add_h(doc, line[2:].strip(), 1)
        elif line.startswith("## "):
            flush_p()
            add_h(doc, line[3:].strip(), 1)
        elif line.startswith("### "):
            flush_p()
            add_h(doc, line[4:].strip(), 2)
        elif line.startswith("#### "):
            flush_p()
            add_h(doc, line[5:].strip(), 3)
        elif line.startswith("> "):
            flush_p()
            body = line[2:]
            kind = "Recomendación"
            for k in ("Confirmado", "Pendiente", "Supuesto de diseño", "Recomendación", "Incidencia"):
                if body.startswith(k):
                    kind = k
                    body = body[len(k) :].lstrip(" .:—-")
                    break
            add_callout(doc, kind, body)
        elif line.startswith("- "):
            flush_p()
            items = []
            while i < len(lines) and lines[i].startswith("- "):
                items.append(lines[i][2:].strip())
                i += 1
            add_bullets(doc, items)
            continue
        elif line.strip() == "":
            flush_p()
        else:
            para_buf.append(line.strip())
        i += 1
    flush_p()
    flush_table()
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(dest)


def write_executive_pdf(dest: Path, kpis: dict):
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(
        str(dest),
        pagesize=A4,
        leftMargin=1.8 * cm,
        rightMargin=1.8 * cm,
        topMargin=1.6 * cm,
        bottomMargin=1.6 * cm,
        title="HC Control · Blueprint ejecutivo v2.0",
        author="HC Control",
    )
    navy = HexColor("#0B1F3A")
    blue = HexColor("#1D4ED8")
    cyan = HexColor("#0E7490")
    ink = HexColor("#0F172A")
    muted = HexColor("#475569")
    ok = HexColor("#065F46")
    styles = getSampleStyleSheet()
    s_h = ParagraphStyle("H", parent=styles["Heading1"], textColor=navy, fontName="Times-Bold", fontSize=18, spaceBefore=14, spaceAfter=8)
    s_h2 = ParagraphStyle("H2", parent=styles["Heading2"], textColor=blue, fontName="Times-Bold", fontSize=13, spaceBefore=10, spaceAfter=6)
    s_b = ParagraphStyle("B", parent=styles["BodyText"], textColor=ink, fontName="Times-Roman", fontSize=10.5, leading=14, spaceAfter=6)
    s_k = ParagraphStyle("K", parent=styles["BodyText"], textColor=white, fontName="Times-Bold", fontSize=9, alignment=1)
    s_m = ParagraphStyle("M", parent=styles["BodyText"], textColor=muted, fontName="Times-Italic", fontSize=9, leading=12)
    s_t = ParagraphStyle("T", parent=styles["Title"], textColor=navy, fontName="Times-Bold", fontSize=26, leading=30, spaceAfter=8)
    story = []
    story.append(Paragraph("HC CONTROL", ParagraphStyle("BR", textColor=blue, fontName="Times-Bold", fontSize=12)))
    story.append(Paragraph("Blueprint ejecutivo v2.0", s_t))
    story.append(Paragraph("Sistema avanzado de control de Headcount · Microsoft 365 · sin Power Apps", s_m))
    story.append(Paragraph("20 de septiembre de 2026 · uso interno · datos de personas sintéticos", s_m))
    story.append(Spacer(1, 10))
    kpi_row = [
        [Paragraph("HC actual", s_k), Paragraph("Incrementos", s_k), Paragraph("Salidas CEV", s_k), Paragraph("Sustituciones", s_k), Paragraph("HC previsto", s_k)],
        [
            Paragraph("514", ParagraphStyle("N", textColor=white, fontName="Times-Bold", fontSize=18, alignment=1)),
            Paragraph("+13", ParagraphStyle("N2", textColor=white, fontName="Times-Bold", fontSize=18, alignment=1)),
            Paragraph("9", ParagraphStyle("N3", textColor=white, fontName="Times-Bold", fontSize=18, alignment=1)),
            Paragraph("9", ParagraphStyle("N4", textColor=white, fontName="Times-Bold", fontSize=18, alignment=1)),
            Paragraph("527", ParagraphStyle("N5", textColor=white, fontName="Times-Bold", fontSize=18, alignment=1)),
        ],
    ]
    t = Table(kpi_row, colWidths=[3.3 * cm] * 5)
    t.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), navy),
                ("BACKGROUND", (4, 0), (4, -1), blue),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ]
        )
    )
    story.append(t)
    story.append(Spacer(1, 8))
    story.append(Paragraph("Ecuación oficial confirmada: 514 + 13 − 9 + 9 = 527. Impacto neto del programa CEV = 0.", s_b))
    story.append(Paragraph("1. Necesidad", s_h))
    story.append(
        Paragraph(
            "La organización necesita una referencia operativa única, fiable, mantenible y trazable del headcount. "
            "No se pretende un sistema de Recursos Humanos: se pretende saber, en todo momento, cuántas personas computan, "
            "cuál es la capacidad estructural autorizada, qué movimientos explican la variación y qué incidencias de calidad impiden un cierre limpio.",
            s_b,
        )
    )
    story.append(Paragraph("2. Arquitectura seleccionada", s_h))
    story.append(
        Paragraph(
            "Nivel A (MVP): una sola lista CONTROL_HC, operable desde Microsoft Lists, Excel y Teams, sin Power BI y sin Power Automate. "
            "Nivel B (objetivo): PERSONAS_HC (quién está), PLAZAS_HC (qué capacidad existe) y MOVIMIENTOS_HC (qué cambia y por qué). "
            "El MVP ya incorpora IDPlaza y MatriculaSustituida para no rehacer el trabajo en la migración.",
            s_b,
        )
    )
    story.append(Paragraph("3. Escenario 514 → 527", s_h))
    evo = [
        ["Mes", "Hecho", "HC"],
        ["Agosto 2026", "Cierre base confirmado", "514"],
        ["Septiembre", "Tras +7 incrementos iniciales (6 becarios + 1 RRLL)", "521"],
        ["Octubre", "+2 posiciones adicionales", "523"],
        ["Noviembre", "+3 posiciones adicionales y CEV pareado (−9/+9)", "526"],
        ["Diciembre", "+1 posición adicional · HC estructural final", "527"],
    ]
    te = Table(evo, colWidths=[3.5 * cm, 10 * cm, 3 * cm])
    te.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), navy),
                ("TEXTCOLOR", (0, 0), (-1, 0), white),
                ("FONTNAME", (0, 0), (-1, 0), "Times-Bold"),
                ("FONTNAME", (0, 1), (-1, -1), "Times-Roman"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("GRID", (0, 0), (-1, -1), 0.3, HexColor("#D0D7E2")),
                ("BACKGROUND", (0, 1), (-1, 1), HexColor("#ECFEFF")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    story.append(te)
    story.append(Paragraph("4. Principios", s_h))
    story.append(
        Paragraph(
            "Las personas explican quién está. Las plazas explican el Headcount. Los movimientos explican qué cambia y por qué. "
            "Incremento crea plaza. Sustitución reutiliza plaza. No se mezclan. Distinción estricta entre previsto y real. "
            "Identificadores como texto. Sin teléfonos ni datos personales innecesarios. Operación por excepción. Cierre mensual con OK/ERROR.",
            s_b,
        )
    )
    story.append(Paragraph("5. Qué está construido", s_h))
    story.append(
        Paragraph(
            "Excel operativo del MVP con tabla CONTROL_HC, catálogos, diccionario, control OK/ERROR y reconciliación 514→527. "
            "CSV de importación a Lists. Plantillas del modelo de tres listas. Especificación de columnas, vistas y JSON de formato. "
            "Manual de construcción por bloques de 15–30 minutos. Especificación Power BI y medidas DAX. Doce flujos Power Automate diseñados "
            "(no operativos hasta configurar el tenant). Plan de pruebas. Roadmap de gobierno. Portal de demostración del paquete.",
            s_b,
        )
    )
    story.append(Paragraph("6. Qué requiere el tenant", s_h))
    story.append(
        Paragraph(
            "Sitio de SharePoint y equipo Teams (DEC-008). Permisos de Lists. Descarga real de RRHH (DEC-010). "
            "Catálogos reales de unidad y empresa (DEC-001, DEC-002). Fechas y referencias de aprobación (DEC-003 a DEC-006). "
            "Workspace Power BI (DEC-013) y cuenta para flujos (DEC-014). Power Apps no se usa.",
            s_b,
        )
    )
    story.append(Paragraph("7. Gobierno mínimo", s_h))
    story.append(
        Paragraph(
            "Operador HC (alta/baja/calidad), Validador (cierre mensual), Lector (dirección). "
            "Automatizaciones solo cuando aportan valor y nunca como requisito del MVP. "
            "Migración a tres listas cuando el MVP esté estable, no antes.",
            s_b,
        )
    )
    story.append(Paragraph("8. Criterio de aceptación del escenario", s_h))
    story.append(
        Paragraph(
            f"HC actual {kpis.get('hc_actual')} · incrementos {kpis.get('incrementos')} · salidas CEV {kpis.get('salidas_cev')} · "
            f"sustituciones {kpis.get('sustituciones')} · HC previsto {kpis.get('hc_previsto_final')} · CEV neto {kpis.get('impacto_cev_neto')} · "
            f"ecuación {'OK' if kpis.get('ecuacion_ok') else 'ERROR'}.",
            s_b,
        )
    )
    story.append(Paragraph("La implantación empieza por el Excel operativo y la lista CONTROL_HC. El resto del paquete está preparado para activarse por fases.", s_b))
    def footer(canvas, doc_):
        canvas.saveState()
        canvas.setFillColor(navy)
        canvas.rect(0, A4[1] - 12, A4[0], 12, fill=1, stroke=0)
        canvas.setFillColor(white)
        canvas.setFont("Times-Roman", 8)
        canvas.drawString(1.8 * cm, A4[1] - 9, "HC Control · Blueprint ejecutivo v2.0")
        canvas.setFillColor(muted)
        canvas.drawRightString(A4[0] - 1.8 * cm, 1.0 * cm, f"Página {doc_.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
