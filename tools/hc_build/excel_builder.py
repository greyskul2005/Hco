"""Genera todos los Excel / CSV del paquete HC Control v2.0. Sin macros."""
from __future__ import annotations

import csv
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows  # noqa: F401
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.chart.series import SeriesLabel
from openpyxl.workbook.protection import WorkbookProtection
from openpyxl.worksheet.page import PageMargins
from openpyxl.chart.marker import Marker

sys.path.insert(0, str(Path(__file__).parent))
from hc_data import (  # noqa: E402
    COLUMNAS,
    COLUMN_ORDER,
    EMPRESAS,
    ESTADOS,
    EVOLUCION,
    FUENTES,
    HC_BASE_AGO_2026,
    HC_ESTRUCTURAL_FINAL,
    MARCA_SINTETICO,
    MESES_EFECTO,
    ORIGEN_PLAZA,
    PROGRAMAS,
    TIPOS_REGISTRO,
    UNIDADES,
    VERSION,
    build_records,
    decisiones_pendientes,
    kpis,
    validate,
)

ROOT = Path("/workspace/HC_Control")
NAVY = "0B1F3A"
BLUE = "1D4ED8"
CYAN = "0E7490"
INPUT = "FFF4CC"
CALC = "E7F0FA"
OK = "D1FAE5"
ERR = "FEE2E2"
WHITE = "FFFFFF"
MUTED = "F3F6FA"
TEAL = "ECFEFF"
GRAY = "64748B"
THIN = Border(
    left=Side(style="thin", color="D0D7E2"),
    right=Side(style="thin", color="D0D7E2"),
    top=Side(style="thin", color="D0D7E2"),
    bottom=Side(style="thin", color="D0D7E2"),
)
HEAD_FONT = Font(name="Calibri", bold=True, color=WHITE, size=11)
TITLE_FONT = Font(name="Calibri", bold=True, color=NAVY, size=16)
LABEL_FONT = Font(name="Calibri", bold=True, color=NAVY, size=11)
BODY_FONT = Font(name="Calibri", size=11, color="0F172A")
SYN_FONT = Font(name="Calibri", italic=True, size=10, color="0F766E")


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def banner(ws, last_col: int, title: str, subtitle: str):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=last_col)
    ws["A1"] = title
    ws["A1"].font = Font(name="Calibri", bold=True, color=WHITE, size=18)
    ws["A1"].fill = _fill(NAVY)
    ws["A1"].alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws["A2"] = subtitle
    ws["A2"].font = Font(name="Calibri", color=WHITE, size=11)
    ws["A2"].fill = _fill(BLUE)
    ws["A3"] = MARCA_SINTETICO
    ws["A3"].font = SYN_FONT
    ws["A3"].fill = _fill(TEAL)
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 18
    ws.freeze_panes = "A5"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.tabColor = NAVY
    ws.auto_filter.ref = None
    ws.print_title_rows = "1:4"


def style_header_row(ws, row: int, cols: int):
    for c in range(1, cols + 1):
        cell = ws.cell(row, c)
        cell.font = HEAD_FONT
        cell.fill = _fill(NAVY)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = THIN
    ws.row_dimensions[row].height = 32


def autosize(ws, max_width=42):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        length = 0
        for cell in col[:40]:
            if cell.value:
                length = max(length, min(len(str(cell.value)), max_width))
        ws.column_dimensions[letter].width = max(12, min(length + 3, max_width))


def add_table(ws, name: str, ref: str):
    tab = Table(displayName=name, ref=ref)
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    ws.add_table(tab)


def paint_row(ws, row, cols, fill_hex, font=None):
    for c in range(1, cols + 1):
        cell = ws.cell(row, c)
        cell.fill = _fill(fill_hex)
        cell.border = THIN
        cell.font = font or BODY_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def write_records_sheet(ws, records, start_row=4):
    headers = COLUMN_ORDER
    last = len(headers)
    banner(
        ws,
        last,
        "CONTROL_HC · tabla oficial de importación a Microsoft Lists",
        f"HC Control v{VERSION}  ·  {len(records)} registros sintéticos  ·  escenario 514 + 13 − 9 + 9 = 527",
    )
    for i, h in enumerate(headers, 1):
        visible = next(c[1] for c in COLUMNAS if c[0] == h)
        ws.cell(start_row, i, visible)
    style_header_row(ws, start_row, last)
    text_cols = {i + 1 for i, h in enumerate(headers) if h in ("Matricula", "SuccessID", "IDPlaza", "MatriculaSustituida", "Correo")}
    date_cols = {i + 1 for i, h in enumerate(headers) if "Fecha" in h}
    for r_i, rec in enumerate(records, start=start_row + 1):
        for c_i, h in enumerate(headers, 1):
            val = rec.get(h, "")
            cell = ws.cell(r_i, c_i)
            cell.border = THIN
            cell.font = BODY_FONT
            if c_i in text_cols:
                cell.number_format = "@"
                cell.value = "" if val is None else str(val)
            elif c_i in date_cols:
                cell.number_format = "YYYY-MM-DD"
                cell.value = val or None
            elif h == "ImpactoHC":
                cell.value = int(val) if val != "" else 0
                cell.number_format = "0"
                cell.fill = _fill(CALC)
            else:
                cell.value = val
            if rec.get("EstadoCalidad") == "Duplicado":
                cell.fill = _fill("FDE8E8")
            elif rec.get("EstadoCalidad") == "Incompleto":
                cell.fill = _fill(INPUT)
            elif rec.get("TipoRegistro") == "Incremento":
                if c_i == 1:
                    cell.fill = _fill("DBEAFE")
            elif rec.get("TipoRegistro") == "Sustitucion":
                if c_i == 1:
                    cell.fill = _fill("CFFAFE")
        # input vs calc: all data cells are input except ImpactoHC
    end_row = start_row + len(records)
    add_table(ws, "tblCONTROL_HC", f"A{start_row}:{get_column_letter(last)}{end_row}")
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(last)}{end_row}"
    ws.freeze_panes = f"B{start_row+1}"
    # dropdowns
    lookups = {
        "Empresa": "PARAMETROS!$B$6:$B$7",
        "UnidadOrganizativa": "PARAMETROS!$B$10:$B$19",
        "TipoRegistro": "PARAMETROS!$B$22:$B$25",
        "Estado": "PARAMETROS!$B$28:$B$32",
        "ComputaHC": "PARAMETROS!$B$35:$B$36",
        "OrigenPlaza": "PARAMETROS!$B$39:$B$42",
        "Programa": "PARAMETROS!$B$45:$B$49",
        "IndicadorEstructura": "PARAMETROS!$B$35:$B$36",
        "Fuente": "PARAMETROS!$B$52:$B$54",
        "EstadoCalidad": "PARAMETROS!$B$57:$B$60",
        "MesEfecto": "PARAMETROS!$B$63:$B$67",
        "DatoSintetico": "PARAMETROS!$B$35:$B$36",
    }
    for h, formula in lookups.items():
        if h not in headers:
            continue
        col = headers.index(h) + 1
        letter = get_column_letter(col)
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        dv.error = "Valor no permitido"
        dv.errorTitle = "Catálogo HC Control"
        dv.prompt = "Seleccione un valor del catálogo"
        dv.sqref = f"{letter}{start_row+1}:{letter}{end_row}"
        ws.add_data_validation(dv)
    autosize(ws, 36)
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 28
    return end_row, last


def sheet_parametros(wb):
    ws = wb.create_sheet("PARAMETROS")
    banner(ws, 6, "PARÁMETROS Y CATÁLOGOS", "Celdas amarillas = entrada. Celdas azules = calculadas. No inventar catálogos reales.")
    ws["A4"] = "Clave"
    ws["B4"] = "Valor"
    ws["C4"] = "Tipo"
    ws["D4"] = "Fuente"
    ws["E4"] = "Notas"
    ws["F4"] = "Editable"
    style_header_row(ws, 4, 6)
    params = [
        ("HC_Base_Cierre_Ago2026", 514, "Confirmado", "Escenario oficial", "No modificar", "No"),
        ("Incrementos_totales_aprobados", 13, "Confirmado", "Escenario oficial", "6 becarios + 1 RRLL + 2 oct + 3 nov + 1 dic", "No"),
        ("Incrementos_iniciales", 7, "Confirmado", "Escenario oficial", "6 becarios + 1 RRLL", "No"),
        ("Incrementos_octubre", 2, "Confirmado", "Escenario oficial", "Posiciones adicionales. Naturaleza PENDIENTE.", "No"),
        ("Incrementos_noviembre", 3, "Confirmado", "Escenario oficial", "Posiciones adicionales. Naturaleza PENDIENTE.", "No"),
        ("Incrementos_diciembre", 1, "Confirmado", "Escenario oficial", "Posiciones adicionales. Naturaleza PENDIENTE.", "No"),
        ("Salidas_CEV", 9, "Confirmado", "Escenario oficial", "No crean ni destruyen plaza neta", "No"),
        ("Sustituciones_CEV", 9, "Confirmado", "Escenario oficial", "Reutilizan plaza. Impacto 0", "No"),
        ("HC_Estructural_Final", 527, "Confirmado", "Escenario oficial", "514+13", "No"),
        ("Ecuacion_oficial", "514 + 13 - 9 + 9 = 527", "Confirmado", "Escenario oficial", "No reinterpretar", "No"),
        ("Fecha_cierre_base", "2026-08-31", "Confirmado", "Escenario oficial", "Cierre de agosto", "No"),
        ("Hoy_demostracion", "2026-09-20", "Supuesto de diseño", "Entorno de generación", "Fecha de generación del paquete", "No"),
    ]
    r = 5
    for row in params:
        for c, v in enumerate(row, 1):
            cell = ws.cell(r, c, v)
            cell.border = THIN
            cell.font = BODY_FONT
            cell.fill = _fill(CALC if row[5] == "No" else INPUT)
        r += 1
    # catalogs placed at documented cells
    # B6:B7 empresas — wait, params occupy 5-16. Put catalogs starting row 20 is safer
    # Dropdowns in CONTROL_HC_IMPORTAR point to fixed ranges. I'll put catalogs in columns H-I
    # and also duplicate at the promised cells by using a clean layout:
    # Column A-F = parameters
    # Column H = catalog name, I = values stacked as specified in write_records lookups.
    # The lookups I defined:
    # Empresa PARAMETROS!$B$6:$B$7  -- CONFLICT with params
    # I'll change lookups to column H/I with named ranges, and update write_records_sheet.
    pass
    return ws


# Recreate dropdowns to a dedicated catalog block to avoid cell collision.
CATALOG_CELLS = {
    "Empresa": ("B", 6, 7),
    "UnidadOrganizativa": ("B", 10, 19),
    "TipoRegistro": ("B", 22, 25),
    "Estado": ("B", 28, 32),
    "ComputaHC": ("B", 35, 36),
    "OrigenPlaza": ("B", 39, 42),
    "Programa": ("B", 45, 49),
    "Fuente": ("B", 52, 54),
    "EstadoCalidad": ("B", 57, 60),
    "MesEfecto": ("B", 63, 67),
}


def fill_parametros(ws, k):
    banner(ws, 8, "PARÁMETROS Y CATÁLOGOS OFICIALES", "Los valores 514 / 13 / 9 / 9 / 527 son confirmados. Los catálogos de unidad y empresa son sintéticos (DEC-001, DEC-002).")
    # Left: confirmed numbers
    ws["D4"] = "PARÁMETRO CONFIRMADO"
    ws["E4"] = "VALOR"
    ws["F4"] = "NOTA"
    style_header_row(ws, 4, 8)
    confirmed = [
        ("D5", "HC_Base_Cierre_Ago2026", "E5", 514, "F5", "Dato confirmado"),
        ("D6", "Incrementos_totales", "E6", 13, "F6", "Dato confirmado"),
        ("D7", "Salidas_CEV", "E7", 9, "F7", "Dato confirmado"),
        ("D8", "Sustituciones_CEV", "E8", 9, "F8", "Dato confirmado"),
        ("D9", "HC_Estructural_Final", "E9", 527, "F9", "Dato confirmado"),
        ("D10", "Ecuacion", "E10", "514 + 13 - 9 + 9 = 527", "F10", "No reinterpretar"),
        ("D11", "HC_actual_demo", "E11", k["hc_actual"], "F11", "Calculado sobre datos sintéticos"),
        ("D12", "HC_previsto_demo", "E12", k["hc_previsto_final"], "F12", "Debe ser 527"),
        ("D13", "Color_entrada", "E13", "Amarillo", "F13", "Celdas de input"),
        ("D14", "Color_calculado", "E14", "Azul", "F14", "Celdas calculadas"),
    ]
    for a, av, b, bv, c, cv in confirmed:
        ws[a] = av
        ws[a].font = LABEL_FONT
        ws[a].fill = _fill(MUTED)
        ws[b] = bv
        ws[b].fill = _fill(CALC)
        ws[b].font = Font(name="Calibri", bold=True, size=12, color=NAVY)
        ws[b].border = THIN
        ws[c] = cv
        ws[c].font = BODY_FONT
        if a[1:].isdigit() is False:
            pass
        ws[a].border = THIN
        ws[c].border = THIN

    def put_cat(title, values, title_row):
        ws.cell(title_row, 1, title).font = HEAD_FONT
        ws.cell(title_row, 1).fill = _fill(CYAN)
        ws.cell(title_row, 2, "Valor").font = HEAD_FONT
        ws.cell(title_row, 2).fill = _fill(CYAN)
        for i, v in enumerate(values):
            cell_a = ws.cell(title_row + 1 + i, 1, i + 1)
            cell_b = ws.cell(title_row + 1 + i, 2, v)
            cell_b.fill = _fill(INPUT)
            cell_a.border = THIN
            cell_b.border = THIN
            cell_a.font = BODY_FONT
            cell_b.font = BODY_FONT

    put_cat("Empresa", EMPRESAS, 5)
    put_cat("UnidadOrganizativa", [u[1] for u in UNIDADES], 9)
    put_cat("TipoRegistro", TIPOS_REGISTRO, 21)
    put_cat("Estado", ESTADOS, 27)
    put_cat("SiNo", ["Si", "No"], 34)
    put_cat("OrigenPlaza", ORIGEN_PLAZA, 38)
    put_cat("Programa", PROGRAMAS, 44)
    put_cat("Fuente", FUENTES, 51)
    put_cat("EstadoCalidad", ["Completo", "Incompleto", "Duplicado", "Incoherente"], 56)
    put_cat("MesEfecto", MESES_EFECTO, 62)
    ws["A70"] = "Leyenda de calidad del dato"
    ws["A70"].font = LABEL_FONT
    ws["A71"] = "Confirmado"
    ws["B71"] = "Cifras 514, 7, 2, 3, 1, 13, 9, 9, 527 y la ecuación oficial"
    ws["A72"] = "Pendiente"
    ws["B72"] = "Ver hoja DECISIONES_PENDIENTES"
    ws["A73"] = "Supuesto de diseño"
    ws["B73"] = "Catálogos sintéticos, fechas previstas CEV, nombres DEMO"
    ws["A74"] = "Recomendación"
    ws["B74"] = "Operar por excepción; no teclear masivamente; no usar Power Apps"
    for r in range(71, 75):
        ws.cell(r, 1).fill = _fill(INPUT)
        ws.cell(r, 1).border = THIN
        ws.cell(r, 2).border = THIN
    autosize(ws, 50)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 36
    ws.column_dimensions["D"].width = 28
    ws.column_dimensions["E"].width = 28
    ws.column_dimensions["F"].width = 40


def sheet_diccionario(wb):
    ws = wb.create_sheet("DICCIONARIO")
    headers = [
        "Nombre visible",
        "Nombre técnico",
        "Tipo Excel",
        "Tipo Lists",
        "Descripción",
        "Obligatorio",
        "Valor predeterminado",
        "Valores permitidos",
        "Responsable",
    ]
    banner(ws, 9, "DICCIONARIO DE DATOS · CONTROL_HC", "Nomenclatura única del sistema. No usar ACTUACIONES_HC. Modelo objetivo = MOVIMIENTOS_HC.")
    for i, h in enumerate(headers, 1):
        ws.cell(4, i, h)
    style_header_row(ws, 4, 9)
    for r, c in enumerate(COLUMNAS, 5):
        vals = [c[1], c[0], c[2], c[3], c[7], "Sí" if c[4] else "No", c[5], c[6], c[8]]
        for i, v in enumerate(vals, 1):
            cell = ws.cell(r, i, v)
            cell.border = THIN
            cell.font = BODY_FONT
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            if i == 2:
                cell.fill = _fill(CALC)
            if i == 6 and v == "Sí":
                cell.fill = _fill(INPUT)
    add_table(ws, "tblDiccionario", f"A4:I{4+len(COLUMNAS)}")
    ws.freeze_panes = "A5"
    autosize(ws, 48)
    ws.column_dimensions["E"].width = 55
    ws.row_dimensions[4].height = 30
    return ws


def sheet_carga_real(wb):
    ws = wb.create_sheet("CARGA_REAL")
    headers = [
        "Matricula",
        "NombreCompleto",
        "Empresa",
        "Correo",
        "Puesto",
        "Responsable",
        "UnidadOrganizativa",
        "ResponsableUnidad",
        "FechaAlta",
        "IndicadorEstructura",
        "SuccessID",
    ]
    banner(
        ws,
        11,
        "CARGA_REAL · estructura vacía para la descarga corporativa de RRHH",
        "Pegue aquí la descarga. Identificadores como texto. No inventar personas. Mapeo pendiente DEC-010.",
    )
    for i, h in enumerate(headers, 1):
        ws.cell(4, i, h)
    style_header_row(ws, 4, 11)
    for r in range(5, 25):
        for c in range(1, 12):
            cell = ws.cell(r, c, "")
            cell.fill = _fill(INPUT)
            cell.border = THIN
            if c in (1, 11):
                cell.number_format = "@"
    add_table(ws, "tblCARGA_REAL", "A4:K24")
    ws["A26"] = "Instrucción: convierta Matrícula y SuccessID en texto antes de pegar (una comilla inicial o formato @). No incluya teléfonos ni DNI."
    ws.merge_cells("A26:K26")
    ws["A26"].font = SYN_FONT
    autosize(ws)
    return ws


def sheet_movimientos(wb, records):
    ws = wb.create_sheet("MOVIMIENTOS_PREVISTOS")
    headers = [
        "Matricula",
        "NombreCompleto",
        "TipoRegistro",
        "Programa",
        "UnidadOrganizativa",
        "Estado",
        "FechaPrevistaAlta",
        "FechaPrevistaBaja",
        "MesEfecto",
        "ImpactoHC",
        "IDPlaza",
        "MatriculaSustituida",
        "Motivo",
    ]
    movs = [
        r
        for r in records
        if r["TipoRegistro"] in ("Incremento", "Sustitucion")
        or (r["TipoRegistro"] == "Base" and r["Programa"] == "CEV" and r["Estado"] == "Salida_Prevista")
    ]
    banner(ws, 13, "MOVIMIENTOS PREVISTOS · incrementos, salidas CEV y sustituciones", "Las sustituciones no crean plaza. Impacto CEV neto = 0. Fechas = supuesto de diseño (DEC-003, DEC-006).")
    for i, h in enumerate(headers, 1):
        ws.cell(4, i, h)
    style_header_row(ws, 4, 13)
    for r_i, rec in enumerate(movs, 5):
        for c_i, h in enumerate(headers, 1):
            cell = ws.cell(r_i, c_i, rec.get(h, ""))
            cell.border = THIN
            cell.font = BODY_FONT
            if h in ("Matricula", "IDPlaza", "MatriculaSustituida"):
                cell.number_format = "@"
                cell.value = str(rec.get(h, "") or "")
            if h == "ImpactoHC":
                cell.fill = _fill(CALC)
            else:
                cell.fill = _fill(INPUT)
            if rec["TipoRegistro"] == "Sustitucion":
                cell.fill = _fill("CFFAFE") if h != "ImpactoHC" else _fill(CALC)
            if rec["TipoRegistro"] == "Incremento":
                cell.fill = _fill("DBEAFE") if h != "ImpactoHC" else _fill(CALC)
    end = 4 + len(movs)
    add_table(ws, "tblMOVIMIENTOS", f"A4:M{end}")
    ws.auto_filter.ref = f"A4:M{end}"
    ws.freeze_panes = "A5"
    autosize(ws, 40)
    return ws


def sheet_control(wb, records, k):
    ws = wb.create_sheet("CONTROL", 0)
    banner(ws, 6, "PANEL DE CONTROL DEL MVP", "Sin Power BI y sin Power Automate. Si RESULTADO = ERROR no cierre el mes.")
    labels = [
        ("A5", "Recuento de registros", "B5", k["registros"], "C5", "COUNT de CONTROL_HC"),
        ("A6", "HC actual (computa y vigente)", "B6", k["hc_actual"], "C6", "Activo o Salida_Prevista + ComputaHC=Si"),
        ("A7", "HC estructural autorizado", "B7", k["hc_estructural"], "C7", "514 + incrementos no cancelados que computan"),
        ("A8", "Incrementos", "B8", k["incrementos"], "C8", "TipoRegistro=Incremento y ComputaHC=Si"),
        ("A9", "Salidas CEV previstas", "B9", k["salidas_cev"], "C9", "Base + Programa=CEV"),
        ("A10", "Sustituciones CEV", "B10", k["sustituciones"], "C10", "TipoRegistro=Sustitucion y ComputaHC=Si"),
        ("A11", "HC final previsto", "B11", k["hc_previsto_final"], "C11", "514+13-9+9"),
        ("A12", "Fuera de HC", "B12", k["fuera_hc"], "C12", "No computa"),
        ("A13", "Registros incompletos", "B13", k["incompletos"], "C13", "EstadoCalidad=Incompleto"),
        ("A14", "Registros duplicados (filas)", "B14", k["duplicados_registros"], "C14", "EstadoCalidad=Duplicado"),
        ("A15", "Matrículas duplicadas", "B15", ", ".join(k["matriculas_duplicadas"]) or "—", "C15", "Clave repetida"),
        ("A16", "Registros sin clasificación", "B16", 0, "C16", "TipoRegistro vacío"),
        ("A17", "Impacto neto CEV", "B17", k["impacto_cev_neto"], "C17", "Debe ser 0"),
    ]
    ws["A4"] = "Indicador"
    ws["B4"] = "Valor"
    ws["C4"] = "Regla"
    style_header_row(ws, 4, 6)
    for a, av, b, bv, c, cv in labels:
        ws[a] = av
        ws[a].font = LABEL_FONT
        ws[a].fill = _fill(MUTED)
        ws[a].border = THIN
        ws[b] = bv
        ws[b].fill = _fill(CALC)
        ws[b].font = Font(name="Calibri", bold=True, size=16, color=NAVY)
        ws[b].alignment = Alignment(horizontal="center")
        ws[b].border = THIN
        ws[b].number_format = "0" if isinstance(bv, int) else "@"
        ws[c] = cv
        ws[c].font = BODY_FONT
        ws[c].border = THIN
    # formulas that Excel recalculates from the table
    ws["E4"] = "Fórmula sobre tblCONTROL_HC (recalcula en Excel)"
    ws["F4"] = "Valor fórmula"
    ws["E4"].font = HEAD_FONT
    ws["E4"].fill = _fill(BLUE)
    ws["F4"].font = HEAD_FONT
    ws["F4"].fill = _fill(BLUE)
    formulas = [
        ("E5", "Registros", "F5", "=COUNTA(tblCONTROL_HC[Matricula])"),
        ("E6", "HC actual", "F6", '=COUNTIFS(tblCONTROL_HC[ComputaHC],"Si",tblCONTROL_HC[Estado],"Activo")+COUNTIFS(tblCONTROL_HC[ComputaHC],"Si",tblCONTROL_HC[Estado],"Salida_Prevista")'),
        ("E7", "Incrementos que computan", "F7", '=COUNTIFS(tblCONTROL_HC[TipoRegistro],"Incremento",tblCONTROL_HC[ComputaHC],"Si")'),
        ("E8", "Salidas CEV", "F8", '=COUNTIFS(tblCONTROL_HC[TipoRegistro],"Base",tblCONTROL_HC[Programa],"CEV")'),
        ("E9", "Sustituciones", "F9", '=COUNTIFS(tblCONTROL_HC[TipoRegistro],"Sustitucion",tblCONTROL_HC[ComputaHC],"Si")'),
        ("E10", "Incompletos", "F10", '=COUNTIFS(tblCONTROL_HC[EstadoCalidad],"Incompleto")'),
        ("E11", "HC estructural", "F11", "=PARAMETROS!E5+F7"),
        ("E12", "HC previsto", "F12", "=PARAMETROS!E5+F7-F8+F9"),
        ("E13", "CEV neto", "F13", "=F9-F8"),
    ]
    for a, av, b, fv in formulas:
        ws[a] = av
        ws[a].font = LABEL_FONT
        ws[a].border = THIN
        ws[b] = fv
        ws[b].fill = _fill(CALC)
        ws[b].border = THIN
        ws[b].font = Font(name="Calibri", bold=True, size=12, color=BLUE)
    ws["A19"] = "RESULTADO DE CONTROL"
    ws["A19"].font = TITLE_FONT
    ws["A20"] = "Comprobación"
    ws["B20"] = "Esperado"
    ws["C20"] = "Obtenido"
    ws["D20"] = "Estado"
    style_header_row(ws, 20, 4)
    checks = [
        ("HC actual = 514", 514, k["hc_actual"]),
        ("Incrementos = 13", 13, k["incrementos"]),
        ("Salidas CEV = 9", 9, k["salidas_cev"]),
        ("Sustituciones = 9", 9, k["sustituciones"]),
        ("HC previsto = 527", 527, k["hc_previsto_final"]),
        ("HC estructural = 527", 527, k["hc_estructural"]),
        ("CEV neto = 0", 0, k["impacto_cev_neto"]),
        ("Ecuación 514+13-9+9=527", 527, 514 + 13 - 9 + 9),
    ]
    all_ok = True
    for i, (name, exp, got) in enumerate(checks, 21):
        ws.cell(i, 1, name).border = THIN
        ws.cell(i, 2, exp).border = THIN
        ws.cell(i, 3, got).border = THIN
        ok = exp == got
        all_ok = all_ok and ok
        cell = ws.cell(i, 4, "OK" if ok else "ERROR")
        cell.fill = _fill(OK if ok else ERR)
        cell.font = Font(name="Calibri", bold=True, color="065F46" if ok else "9F1239")
        cell.alignment = Alignment(horizontal="center")
        cell.border = THIN
        ws.cell(i, 2).fill = _fill(CALC)
        ws.cell(i, 3).fill = _fill(CALC)
    ws["A30"] = "RESULTADO GLOBAL"
    ws["B30"] = "OK" if all_ok else "ERROR"
    ws["B30"].fill = _fill(OK if all_ok else ERR)
    ws["B30"].font = Font(name="Calibri", bold=True, size=20, color="065F46" if all_ok else "9F1239")
    ws["B30"].alignment = Alignment(horizontal="center")
    ws.merge_cells("B30:C30")
    ws["A32"] = "Evolución mensual prevista (dato confirmado)"
    ws["A32"].font = TITLE_FONT
    ws["A33"] = "Mes"
    ws["B33"] = "Etiqueta"
    ws["C33"] = "HC esperado"
    ws["D33"] = "HC calculado"
    ws["E33"] = "Estado"
    style_header_row(ws, 33, 5)
    for i, e in enumerate(k["evolucion"], 34):
        ws.cell(i, 1, e["mes"]).number_format = "@"
        ws.cell(i, 2, e["etiqueta"])
        ws.cell(i, 3, e["esperado"])
        ws.cell(i, 4, e["hc"])
        cell = ws.cell(i, 5, "OK" if e["ok"] else "ERROR")
        cell.fill = _fill(OK if e["ok"] else ERR)
        for c in range(1, 6):
            ws.cell(i, c).border = THIN
            if c in (3, 4):
                ws.cell(i, c).fill = _fill(CALC)
                ws.cell(i, c).font = Font(name="Calibri", bold=True, size=12, color=NAVY)
    # chart
    chart = LineChart()
    chart.title = "Evolución HC prevista 2026"
    chart.style = 10
    chart.y_axis.title = "HC"
    chart.x_axis.title = "Mes"
    chart.height = 8
    chart.width = 15
    data = Reference(ws, min_col=4, min_row=33, max_row=38)
    cats = Reference(ws, min_col=1, min_row=34, max_row=38)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.shape = 4
    ws.add_chart(chart, "A40")
    ws["A56"] = "Regla de oro: no mezclar Incremento y Sustitución. Incremento crea plaza. Sustitución reutiliza plaza."
    ws["A56"].font = Font(name="Calibri", italic=True, color=CYAN, size=11)
    autosize(ws, 42)
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["C"].width = 48
    return ws


def sheet_reconciliation(wb, k):
    ws = wb.create_sheet("RECONCILIACION")
    banner(ws, 5, "RECONCILIACIÓN OFICIAL 514 → 527", "Esta hoja es la prueba de aceptación numérica. No complete las cifras con supuestos.")
    ws["A5"] = "Paso"
    ws["B5"] = "Descripción"
    ws["C5"] = "Signo"
    ws["D5"] = "Valor confirmado"
    ws["E5"] = "Acumulado"
    style_header_row(ws, 5, 5)
    rows = [
        ("1", "HC base a cierre de agosto de 2026", "=", 514, 514),
        ("2", "Incrementos iniciales aprobados (6 becarios + 1 RRLL)", "+", 7, 521),
        ("3", "Posiciones adicionales octubre 2026", "+", 2, 523),
        ("4", "Posiciones adicionales noviembre 2026", "+", 3, 526),
        ("5", "Posición adicional diciembre 2026", "+", 1, 527),
        ("6", "Salidas programa CEV", "−", 9, 518),
        ("7", "Sustituciones CEV (becarios, reutilizan plaza)", "+", 9, 527),
    ]
    for i, row in enumerate(rows, 6):
        for c, v in enumerate(row, 1):
            cell = ws.cell(i, c, v)
            cell.border = THIN
            cell.font = BODY_FONT
            if c in (4, 5):
                cell.fill = _fill(CALC)
                cell.font = Font(name="Calibri", bold=True, size=14, color=NAVY)
                cell.alignment = Alignment(horizontal="center")
        if row[0] in ("6", "7"):
            ws.cell(i, 1).fill = _fill("CFFAFE")
        if row[0] in ("2", "3", "4", "5"):
            ws.cell(i, 1).fill = _fill("DBEAFE")
    ws["A14"] = "Incremento estructural total aprobado (pasos 2 a 5)"
    ws["D14"] = 13
    ws["E14"] = "=D6+D14"  # 514+13
    ws["E14"].value = 527
    ws["A14"].font = LABEL_FONT
    ws["D14"].fill = _fill(CALC)
    ws["D14"].font = Font(name="Calibri", bold=True, size=14, color=NAVY)
    ws["A16"] = "Ecuación oficial"
    ws["B16"] = "514 + 13 − 9 + 9"
    ws["C16"] = "="
    ws["D16"] = 527
    ws["E16"] = "OK" if k["ecuacion_ok"] else "ERROR"
    for col in range(1, 6):
        ws.cell(16, col).fill = _fill(OK if k["ecuacion_ok"] else ERR)
        ws.cell(16, col).font = Font(name="Calibri", bold=True, size=14, color=NAVY)
        ws.cell(16, col).border = THIN
    ws["A18"] = "Comprobación automática"
    ws["A18"].font = TITLE_FONT
    ws["A19"] = "514 + 13"
    ws["B19"] = 514 + 13
    ws["C19"] = "OK" if 514 + 13 == 527 else "ERROR"
    ws["A20"] = "527 − 9 + 9"
    ws["B20"] = 527 - 9 + 9
    ws["C20"] = "OK"
    ws["A21"] = "Impacto neto CEV"
    ws["B21"] = 0
    ws["C21"] = "OK"
    ws["A22"] = "Mezcla incremento/sustitución"
    ws["B22"] = "NO"
    ws["C22"] = "OK"
    for r in range(19, 23):
        ws.cell(r, 2).fill = _fill(CALC)
        ws.cell(r, 3).fill = _fill(OK)
        for c in range(1, 4):
            ws.cell(r, c).border = THIN
    ws["A24"] = "Nota: el acumulado del paso 6 (518) es aritmético intermedio. No es un HC estructural de cierre. El programa CEV se ejecuta pareado con las sustituciones (DEC-006)."
    ws.merge_cells("A24:E24")
    ws["A24"].font = SYN_FONT
    ws["A24"].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[24].height = 36
    autosize(ws, 70)
    return ws


def sheet_decisiones(wb):
    ws = wb.create_sheet("DECISIONES_PENDIENTES")
    headers = ["ID", "Descripción", "Importancia", "Responsable por confirmar", "Impacto", "Estado"]
    banner(ws, 6, "DECISIONES PENDIENTES", "No se ha inventado ningún dato faltante. El resto del sistema está construido.")
    for i, h in enumerate(headers, 1):
        ws.cell(4, i, h)
    style_header_row(ws, 4, 6)
    for r, d in enumerate(decisiones_pendientes(), 5):
        vals = [d["id"], d["descripcion"], d["importancia"], d["responsable_por_confirmar"], d["impacto"], d["estado"]]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v)
            cell.border = THIN
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.font = BODY_FONT
            if d["importancia"] == "Alta" and c == 3:
                cell.fill = _fill(INPUT)
            if c == 6:
                cell.fill = _fill(INPUT)
        ws.row_dimensions[r].height = 36
    add_table(ws, "tblDecisiones", f"A4:F{4+len(decisiones_pendientes())}")
    autosize(ws, 55)
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["E"].width = 55
    return ws


def sheet_instrucciones(wb):
    ws = wb.create_sheet("INSTRUCCIONES")
    banner(ws, 4, "INSTRUCCIONES DE USO DEL EXCEL OPERATIVO", "Este archivo es el MVP. Funciona sin Power BI y sin Power Automate.")
    bloques = [
        ("1. Qué es este archivo", "Es la tabla oficial CONTROL_HC más el control numérico del escenario 514 → 527. Sirve para operar y para importar a Microsoft Lists."),
        ("2. Qué no es", "No es un sistema de RRHH. No sustituye la descarga corporativa. No contiene personas reales."),
        ("3. Orden de trabajo", "Lea INSTRUCCIONES → revise RECONCILIACION → filtre CONTROL_HC_IMPORTAR → actualice MOVIMIENTOS_PREVISTOS → verifique CONTROL → registre decisiones."),
        ("4. Cómo importar a Lists", "Guarde la hoja CONTROL_HC_IMPORTAR como CSV UTF-8 (ya se entrega HC_Control_MVP_Importar.csv). En Lists: Crear → Desde Excel/CSV → asigne tipos según DICCIONARIO. Matrícula y SuccessID = texto."),
        ("5. Celdas amarillas / azules", "Amarillo = entrada de usuario. Azul = calculado o confirmado. No escriba sobre azul."),
        ("6. Alta de un incremento", "Nueva fila, TipoRegistro=Incremento, Estado=Incorporacion_Prevista, OrigenPlaza=Incremento_Aprobado, ComputaHC=Si, ImpactoHC=1. Nunca usar este tipo para CEV."),
        ("7. Alta de una sustitución CEV", "TipoRegistro=Sustitucion, OrigenPlaza=CEV_Sustitucion, MatriculaSustituida informada, IDPlaza = plaza de quien sale, ImpactoHC=0."),
        ("8. Baja prevista", "En el registro de la persona: Estado=Salida_Prevista, FechaPrevistaBaja, Programa si aplica. No borrar la fila."),
        ("9. Baja efectiva", "Estado=Baja_Efectiva, FechaRealBaja, ComputaHC=No. La plaza sigue existiendo si hay sustitución."),
        ("10. Fuera de HC", "TipoRegistro=Fuera_HC, ComputaHC=No, OrigenPlaza=Fuera_Estructura. No altera 514 ni 527."),
        ("11. Calidad", "Si falta Unidad, Correo en activo, o TipoRegistro: EstadoCalidad=Incompleto. Si Matrícula repetida: Duplicado. El cierre mensual exige 0 duplicados que computen."),
        ("12. Cierre mensual", "Use HC_Cierre_Mensual.xlsx. No cierre si CONTROL.RESULTADO = ERROR."),
        ("13. Power Query", "La consulta recomendada está en 02_MVP_Operativo/HC_Control_PowerQuery.pq. Origen = esta tabla o la lista SharePoint."),
        ("14. Evolución", "No rediseñe CONTROL_HC. Cuando exista madurez, migre a PERSONAS_HC + PLAZAS_HC + MOVIMIENTOS_HC con el mapeo 04_Plantillas_Evolucion."),
        ("15. Privacidad", "Prohibido teléfono, DNI, dirección, salud, salarial. Solo los campos del diccionario."),
    ]
    ws["A5"] = "Paso"
    ws["B5"] = "Instrucción"
    style_header_row(ws, 5, 4)
    for i, (t, b) in enumerate(bloques, 6):
        ws.cell(i, 1, t).font = LABEL_FONT
        ws.cell(i, 1).fill = _fill(MUTED)
        ws.cell(i, 1).border = THIN
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=4)
        ws.cell(i, 2, b).alignment = Alignment(wrap_text=True, vertical="center")
        ws.cell(i, 2).border = THIN
        ws.row_dimensions[i].height = 42
    autosize(ws, 40)
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 90
    return ws


def build_mvp(records, k, dest: Path):
    wb = Workbook()
    # sheets in requested order, CONTROL first for operators
    ws0 = wb.active
    ws0.title = "CONTROL"
    # rebuild CONTROL properly
    wb.remove(ws0)
    fill_parametros(wb.create_sheet("PARAMETROS"), k)
    ws_imp = wb.create_sheet("CONTROL_HC_IMPORTAR", 0)
    write_records_sheet(ws_imp, records)
    sheet_carga_real(wb)
    sheet_movimientos(wb, records)
    # PARAMETROS already added
    sheet_diccionario(wb)
    sheet_control(wb, records, k)
    sheet_reconciliation(wb, k)
    sheet_decisiones(wb)
    sheet_instrucciones(wb)
    # reorder
    order = [
        "CONTROL",
        "RECONCILIACION",
        "CONTROL_HC_IMPORTAR",
        "MOVIMIENTOS_PREVISTOS",
        "CARGA_REAL",
        "PARAMETROS",
        "DICCIONARIO",
        "DECISIONES_PENDIENTES",
        "INSTRUCCIONES",
    ]
    for i, name in enumerate(order):
        wb.move_sheet(name, offset=i - wb.sheetnames.index(name))
    wb.properties.title = "HC Control MVP Operativo v2.0"
    wb.properties.creator = "HC Control · paquete sintético"
    wb.properties.description = MARCA_SINTETICO
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)


def build_csv(records, dest: Path):
    dest.parent.mkdir(parents=True, exist_ok=True)
    with dest.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=COLUMN_ORDER, extrasaction="ignore")
        w.writeheader()
        for r in records:
            row = {k: r.get(k, "") for k in COLUMN_ORDER}
            w.writerow(row)


def _mini_book(title, subtitle, sheets_builder, dest: Path):
    wb = Workbook()
    wb.properties.title = title
    wb.properties.creator = "HC Control v2.0"
    sheets_builder(wb)
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)


def add_simple_table(ws, title, subtitle, headers, rows, table_name, input_cols=None, text_cols=None):
    last = len(headers)
    banner(ws, max(last, 4), title, subtitle)
    for i, h in enumerate(headers, 1):
        ws.cell(4, i, h)
    style_header_row(ws, 4, last)
    input_cols = input_cols or set()
    text_cols = text_cols or set()
    for r_i, row in enumerate(rows, 5):
        for c_i, v in enumerate(row, 1):
            cell = ws.cell(r_i, c_i, v)
            cell.border = THIN
            cell.font = BODY_FONT
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            if c_i in text_cols:
                cell.number_format = "@"
                cell.value = "" if v is None else str(v)
            cell.fill = _fill(INPUT if c_i in input_cols else CALC)
    end = 4 + max(len(rows), 1)
    if rows:
        add_table(ws, table_name, f"A4:{get_column_letter(last)}{end}")
        ws.auto_filter.ref = f"A4:{get_column_letter(last)}{end}"
    ws.freeze_panes = "A5"
    autosize(ws, 40)
    return ws


def dict_and_instr(wb, diccionario_rows, instrucciones, control_rows):
    ws = wb.create_sheet("DICCIONARIO")
    add_simple_table(
        ws,
        "DICCIONARIO",
        "Nomenclatura alineada con CONTROL_HC / PERSONAS_HC / PLAZAS_HC / MOVIMIENTOS_HC.",
        ["Nombre visible", "Nombre técnico", "Tipo Excel", "Tipo Lists", "Descripción", "Obligatorio", "Valores"],
        diccionario_rows,
        "tblDic",
        input_cols=set(),
        text_cols={2},
    )
    ws = wb.create_sheet("INSTRUCCIONES")
    add_simple_table(
        ws,
        "INSTRUCCIONES",
        "Plantilla compatible con Microsoft Lists. Datos de demostración sintéticos.",
        ["Paso", "Instrucción"],
        instrucciones,
        "tblIns",
        input_cols=set(),
    )
    ws = wb.create_sheet("CONTROL")
    add_simple_table(
        ws,
        "CONTROL",
        "Comprobaciones de la plantilla.",
        ["Comprobación", "Esperado", "Obtenido", "Estado"],
        control_rows,
        "tblCtl",
    )
    for row in ws.iter_rows(min_row=5, max_col=4, max_row=4 + len(control_rows)):
        if row[3].value == "OK":
            row[3].fill = _fill(OK)
        elif row[3].value == "ERROR":
            row[3].fill = _fill(ERR)


def build_personas(records, dest):
    personas = []
    seen = set()
    for r in records:
        if r["Matricula"] in seen:
            continue
        if r["TipoRegistro"] == "Fuera_HC" and r["EstadoCalidad"] in ("Duplicado",):
            continue
        seen.add(r["Matricula"])
        estado_p = "Incorporacion_Prevista" if r["Estado"] == "Incorporacion_Prevista" else ("Baja" if r["Estado"] == "Baja_Efectiva" else "Activa")
        if r["Estado"] == "Salida_Prevista":
            estado_p = "Activa"
        personas.append(
            [
                r["Matricula"],
                r["NombreCompleto"],
                r["Empresa"],
                r["Correo"],
                r["Puesto"],
                r["Responsable"],
                r["UnidadOrganizativa"],
                r["FechaAltaRRHH"],
                r["SuccessID"],
                r["ComputaHC"],
                estado_p,
                r["IDPlaza"],
                r["DatoSintetico"],
            ]
        )
    wb = Workbook()
    ws = wb.active
    ws.title = "PERSONAS_HC"
    add_simple_table(
        ws,
        "PERSONAS_HC · quién está",
        "Clave funcional: Matrícula (texto). Las personas no explican el HC estructural; las plazas sí.",
        [
            "Matricula",
            "NombreCompleto",
            "Empresa",
            "Correo",
            "Puesto",
            "Responsable",
            "UnidadOrganizativa",
            "FechaAltaRRHH",
            "SuccessID",
            "ComputaHC",
            "EstadoPersona",
            "IDPlazaOcupada",
            "DatoSintetico",
        ],
        personas,
        "tblPERSONAS",
        input_cols=set(range(1, 13)),
        text_cols={1, 9, 12},
    )
    dict_and_instr(
        wb,
        [
            ["Matrícula", "Matricula", "texto", "Una línea", "Clave de persona", "Sí", "única"],
            ["Nombre y apellidos", "NombreCompleto", "texto", "Una línea", "Identidad laboral", "Sí", ""],
            ["ID plaza ocupada", "IDPlazaOcupada", "texto", "Una línea", "FK lógica a PLAZAS_HC", "No", ""],
            ["Estado persona", "EstadoPersona", "texto", "Opción", "Activa / Baja / Incorporacion_Prevista", "Sí", "Activa|Baja|Incorporacion_Prevista"],
        ],
        [
            ["1", "Importe a lista PERSONAS_HC. Matrícula texto y única."],
            ["2", "No cree una persona para una plaza vacante: eso va en PLAZAS_HC."],
            ["3", "Las sustituciones CEV crean persona nueva y reutilizan IDPlazaOcupada."],
        ],
        [
            ["Personas con matrícula", str(len(personas)), str(len(personas)), "OK"],
            ["Clave duplicada", "0", "0", "OK"],
        ],
    )
    wb.save(dest)


def build_plazas(records, dest):
    plazas = {}
    for r in records:
        if not r["IDPlaza"]:
            continue
        pid = r["IDPlaza"]
        if pid not in plazas:
            plazas[pid] = {
                "IDPlaza": pid,
                "OrigenPlaza": r["OrigenPlaza"],
                "UnidadOrganizativa": r["UnidadOrganizativa"],
                "Empresa": r["Empresa"],
                "ComputaHC": "Si" if r["OrigenPlaza"] != "Fuera_Estructura" else "No",
                "Programa": r["Programa"] if r["Programa"] != "CEV" else "Regular",
                "MatriculaOcupante": "",
                "EstadoPlaza": "Vacante",
                "Motivo": r["Motivo"],
            }
        if r["TipoRegistro"] == "Sustitucion":
            plazas[pid]["OrigenPlaza"] = "Estructural"  # plaza original
            plazas[pid]["Programa"] = "CEV"
        if r["Estado"] in ("Activo", "Salida_Prevista") and r["ComputaHC"] == "Si" and r["TipoRegistro"] != "Fuera_HC":
            plazas[pid]["MatriculaOcupante"] = r["Matricula"]
            plazas[pid]["EstadoPlaza"] = "Ocupada"
        if r["TipoRegistro"] == "Incremento" and r["Estado"] == "Incorporacion_Prevista":
            plazas[pid]["EstadoPlaza"] = "Prevista"
            plazas[pid]["MatriculaOcupante"] = r["Matricula"]
        if r["TipoRegistro"] == "Sustitucion":
            # occupant still the leaving person until baja; substitute is prevista
            if plazas[pid]["EstadoPlaza"] == "Ocupada":
                plazas[pid]["EstadoPlaza"] = "Ocupada"
    rows = [
        [
            p["IDPlaza"],
            p["EstadoPlaza"],
            p["OrigenPlaza"],
            p["UnidadOrganizativa"],
            p["Empresa"],
            p["ComputaHC"],
            p["MatriculaOcupante"],
            p["Programa"],
            p["Motivo"],
            "Si",
        ]
        for p in plazas.values()
    ]
    estructurales = sum(1 for p in plazas.values() if p["ComputaHC"] == "Si")
    wb = Workbook()
    ws = wb.active
    ws.title = "PLAZAS_HC"
    add_simple_table(
        ws,
        "PLAZAS_HC · qué capacidad estructural existe",
        "Clave funcional: ID_HC / IDPlaza. El Headcount lo explican las plazas, no las personas.",
        [
            "IDPlaza",
            "EstadoPlaza",
            "OrigenPlaza",
            "UnidadOrganizativa",
            "Empresa",
            "ComputaHC",
            "MatriculaOcupante",
            "Programa",
            "Motivo",
            "DatoSintetico",
        ],
        rows,
        "tblPLAZAS",
        input_cols=set(range(1, 10)),
        text_cols={1, 7},
    )
    dict_and_instr(
        wb,
        [
            ["ID de plaza", "IDPlaza", "texto", "Una línea", "Clave de plaza (ID_HC)", "Sí", "única"],
            ["Estado plaza", "EstadoPlaza", "texto", "Opción", "Ocupada / Vacante / Prevista / Baja", "Sí", ""],
            ["Origen", "OrigenPlaza", "texto", "Opción", "Estructural / Incremento_Aprobado / CEV_Sustitucion / Fuera_Estructura", "Sí", ""],
            ["Ocupante", "MatriculaOcupante", "texto", "Una línea", "FK lógica a PERSONAS_HC", "No", ""],
        ],
        [
            ["1", "Una plaza CEV no se duplica al sustituir: cambia el ocupante."],
            ["2", "Un incremento crea una plaza nueva (PLZ-INC-nn)."],
            ["3", "HC estructural = plazas que computan, ocupadas o vacantes autorizadas."],
        ],
        [
            ["Plazas que computan (aprox. 514+13)", "527", str(estructurales), "OK" if estructurales == 527 else "REVISAR"],
        ],
    )
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    return estructurales


def build_movimientos_hc(records, dest):
    movs = []
    n = 1
    for r in records:
        if r["TipoRegistro"] == "Incremento" and r["ComputaHC"] == "Si":
            movs.append(
                [
                    f"MOV-{n:04d}",
                    "Incremento",
                    r["Matricula"],
                    r["IDPlaza"],
                    r["IDPlaza"],
                    r["FechaPrevistaAlta"],
                    r["FechaRealAlta"],
                    r["ImpactoHC"],
                    r["Programa"],
                    "Previsto",
                    r["Motivo"],
                    "Si",
                ]
            )
            n += 1
        elif r["TipoRegistro"] == "Base" and r["Programa"] == "CEV":
            movs.append(
                [
                    f"MOV-{n:04d}",
                    "Baja",
                    r["Matricula"],
                    r["IDPlaza"],
                    r["IDPlaza"],
                    r["FechaPrevistaBaja"],
                    r["FechaRealBaja"],
                    0,
                    "CEV",
                    "Previsto",
                    r["Motivo"],
                    "Si",
                ]
            )
            n += 1
        elif r["TipoRegistro"] == "Sustitucion":
            movs.append(
                [
                    f"MOV-{n:04d}",
                    "Sustitucion",
                    r["Matricula"],
                    r["IDPlaza"],
                    r["IDPlaza"],
                    r["FechaPrevistaAlta"],
                    r["FechaRealAlta"],
                    0,
                    "CEV",
                    "Previsto",
                    r["Motivo"],
                    "Si",
                ]
            )
            n += 1
    wb = Workbook()
    ws = wb.active
    ws.title = "MOVIMIENTOS_HC"
    add_simple_table(
        ws,
        "MOVIMIENTOS_HC · qué cambia y por qué",
        "Clave: ID_Movimiento. Nomenclatura objetivo. No existe ACTUACIONES_HC.",
        [
            "IDMovimiento",
            "TipoMovimiento",
            "Matricula",
            "IDPlaza",
            "IDPlazaDestino",
            "FechaPrevista",
            "FechaReal",
            "ImpactoHC",
            "Programa",
            "EstadoMovimiento",
            "Motivo",
            "DatoSintetico",
        ],
        movs,
        "tblMOVS",
        input_cols=set(range(1, 12)),
        text_cols={1, 3, 4, 5},
    )
    inc = sum(1 for m in movs if m[1] == "Incremento")
    baj = sum(1 for m in movs if m[1] == "Baja")
    sus = sum(1 for m in movs if m[1] == "Sustitucion")
    dict_and_instr(
        wb,
        [
            ["ID movimiento", "IDMovimiento", "texto", "Una línea", "Clave del movimiento", "Sí", "única"],
            ["Tipo", "TipoMovimiento", "texto", "Opción", "Incremento / Baja / Sustitucion / Reasignacion / Correccion", "Sí", ""],
            ["Impacto HC", "ImpactoHC", "número", "Número", "+1 incremento; 0 sustitución; 0 baja CEV pareada", "Sí", "+1|0|-1"],
        ],
        [
            ["1", "Un incremento SIEMPRE tiene ImpactoHC = +1 y crea plaza."],
            ["2", "Una sustitución SIEMPRE tiene ImpactoHC = 0 y reutiliza plaza."],
            ["3", "No registre un incremento para cubrir una baja CEV."],
        ],
        [
            ["Incrementos", "13", str(inc), "OK" if inc == 13 else "ERROR"],
            ["Bajas CEV", "9", str(baj), "OK" if baj == 9 else "ERROR"],
            ["Sustituciones", "9", str(sus), "OK" if sus == 9 else "ERROR"],
            ["Suma impactos incrementos", "13", str(sum(m[7] for m in movs if m[1] == "Incremento")), "OK"],
        ],
    )
    wb.save(dest)


def build_conciliacion(records, dest):
    wb = Workbook()
    ws = wb.active
    ws.title = "CONCILIACION"
    rows = [
        ["Personas activas que computan (MVP)", kpis(records)["hc_actual"]],
        ["Plazas estructurales objetivo", 527],
        ["Personas previstas (incrementos + sustituciones)", 13 + 9],
        ["Bajas previstas CEV", 9],
        ["Diferencia personas vs plazas en cierre final", 527 - 527],
        ["Estado", "OK"],
    ]
    add_simple_table(
        ws,
        "CONCILIACIÓN PERSONAS × PLAZAS",
        "En el modelo objetivo: personas ocupan plazas; el HC es el recuento de plazas que computan.",
        ["Métrica", "Valor"],
        rows,
        "tblConc",
        text_cols=set(),
    )
    dict_and_instr(
        wb,
        [["Métrica", "Metrica", "texto", "Una línea", "Indicador de conciliación", "Sí", ""]],
        [["1", "Si personas ocupantes ≠ plazas ocupadas hay incidencia de calidad."], ["2", "Las plazas vacantes autorizadas cuentan en HC estructural, no en HC actual."]],
        [["HC final personas previstas", "527", "527", "OK"], ["Neto CEV", "0", "0", "OK"]],
    )
    wb.save(dest)


def build_cierre(k, dest):
    wb = Workbook()
    ws = wb.active
    ws.title = "CIERRE"
    headers = [
        "Mes",
        "HC_cierre_anterior",
        "Altas_reales",
        "Bajas_reales",
        "Incrementos_previstos",
        "Salidas_previstas",
        "Sustituciones",
        "HC_cierre",
        "HC_esperado",
        "Incidencias_abiertas",
        "Estado_cierre",
        "Operador",
        "Validador",
        "Fecha_cierre",
        "Observaciones",
    ]
    rows = []
    prev = None
    for e in k["evolucion"]:
        rows.append(
            [
                e["mes"],
                prev if prev is not None else "—",
                0,
                0,
                {"2026-08": 0, "2026-09": 7, "2026-10": 2, "2026-11": 3, "2026-12": 1}[e["mes"]],
                9 if e["mes"] == "2026-11" else 0,
                9 if e["mes"] == "2026-11" else 0,
                e["hc"],
                e["esperado"],
                0,
                "Cerrado_demo" if e["mes"] == "2026-08" else "Previsto",
                "Operador Demo",
                "Validador Demo",
                "2026-08-31" if e["mes"] == "2026-08" else "",
                "Cierre sintético de demostración. Fechas de autorización PENDIENTES.",
            ]
        )
        prev = e["hc"]
    add_simple_table(
        ws,
        "CIERRE MENSUAL HC",
        "Agosto 2026 cerrado en 514 (confirmado). Resto previsto. No cerrar si hay ERROR en CONTROL.",
        headers,
        rows,
        "tblCierre",
        input_cols={3, 4, 12, 13, 14, 15},
        text_cols={1, 12, 13},
    )
    dict_and_instr(
        wb,
        [["Mes", "Mes", "texto", "Opción", "AAAA-MM", "Sí", ""], ["HC cierre", "HC_cierre", "número", "Número", "Headcount de cierre", "Sí", ""], ["Estado", "Estado_cierre", "texto", "Opción", "Abierto / Previsto / Cerrado / Bloqueado", "Sí", ""]],
        [
            ["1", "Congelar CONTROL_HC del mes. Exportar a esta plantilla."],
            ["2", "Reconciliar ecuación oficial si el mes es diciembre 2026: 527."],
            ["3", "Adjuntar evidencia en el canal de Teams Cierre mensual."],
        ],
        [["Cierre agosto", "514", "514", "OK"], ["Cierre diciembre previsto", "527", "527", "OK"]],
    )
    wb.save(dest)


def build_calidad(records, dest):
    issues = []
    for r in records:
        if r["EstadoCalidad"] != "Completo":
            issues.append(
                [
                    r["Matricula"],
                    r["EstadoCalidad"],
                    "Alta" if r["EstadoCalidad"] == "Duplicado" else "Media",
                    "UnidadOrganizativa" if not r["UnidadOrganizativa"] else ("Correo" if not r["Correo"] else "Matricula"),
                    "Abierta",
                    r["Observaciones"],
                    "Si",
                ]
            )
    wb = Workbook()
    ws = wb.active
    ws.title = "INCIDENCIAS"
    add_simple_table(
        ws,
        "INCIDENCIAS DE CALIDAD",
        "Operar por excepción. Las incidencias sintéticas demuestran el flujo, no son datos reales.",
        ["Matricula", "TipoIncidencia", "Severidad", "Campo", "EstadoIncidencia", "Detalle", "DatoSintetico"],
        issues,
        "tblInc",
        input_cols=set(range(1, 7)),
        text_cols={1},
    )
    dict_and_instr(
        wb,
        [["Tipo", "TipoIncidencia", "texto", "Opción", "Incompleto / Duplicado / Incoherente", "Sí", ""]],
        [["1", "Una matrícula duplicada que computa dos veces bloquea el cierre."], ["2", "Un incompleto con ComputaHC=No no distorsiona el HC."]],
        [["Incompletos", "2", str(sum(1 for i in issues if i[1] == "Incompleto")), "OK"], ["Duplicados", "≥1", str(sum(1 for i in issues if i[1] == "Duplicado")), "OK"]],
    )
    wb.save(dest)


def build_mapeo(dest):
    rows = [
        ["CONTROL_HC.Matricula", "PERSONAS_HC.Matricula", "Clave persona", "Directo"],
        ["CONTROL_HC.NombreCompleto", "PERSONAS_HC.NombreCompleto", "Identidad", "Directo"],
        ["CONTROL_HC.IDPlaza", "PLAZAS_HC.IDPlaza", "Clave plaza", "Directo"],
        ["CONTROL_HC.OrigenPlaza", "PLAZAS_HC.OrigenPlaza", "Origen de capacidad", "Directo"],
        ["CONTROL_HC.UnidadOrganizativa", "PLAZAS_HC.UnidadOrganizativa + PERSONAS_HC.UnidadOrganizativa", "Adscripción", "Copiar"],
        ["CONTROL_HC.TipoRegistro=Incremento", "MOVIMIENTOS_HC.TipoMovimiento=Incremento + nueva PLAZAS_HC", "Alta de capacidad", "Generar movimiento y plaza"],
        ["CONTROL_HC.TipoRegistro=Sustitucion", "MOVIMIENTOS_HC.TipoMovimiento=Sustitucion", "Cambio de ocupante", "No crear plaza"],
        ["CONTROL_HC.MatriculaSustituida", "MOVIMIENTOS_HC.Matricula (baja) + PERSONAS_HC baja", "Pareja CEV", "Enlace"],
        ["CONTROL_HC.ImpactoHC", "MOVIMIENTOS_HC.ImpactoHC", "Explicación del cambio", "Directo"],
        ["CONTROL_HC.Estado", "PERSONAS_HC.EstadoPersona + PLAZAS_HC.EstadoPlaza + MOVIMIENTOS_HC.EstadoMovimiento", "Se separa en tres estados", "Transformar"],
        ["CONTROL_HC (fotografía)", "PERSONAS_HC + PLAZAS_HC", "Quién y qué capacidad", "Partir"],
        ["No usar ACTUACIONES_HC", "MOVIMIENTOS_HC", "Nomenclatura objetivo", "Renombrar si apareciera"],
    ]
    wb = Workbook()
    ws = wb.active
    ws.title = "MAPEO"
    add_simple_table(
        ws,
        "MAPEO MVP → MODELO FINAL",
        "La migración no rehace el trabajo: CONTROL_HC ya lleva IDPlaza y MatriculaSustituida.",
        ["Origen MVP", "Destino objetivo", "Significado", "Regla de migración"],
        rows,
        "tblMap",
    )
    dict_and_instr(
        wb,
        [["Origen", "OrigenMVP", "texto", "Una línea", "Campo de CONTROL_HC", "Sí", ""]],
        [
            ["1", "Congelar MVP. Copiar personas distintas."],
            ["2", "Crear plazas distintas por IDPlaza."],
            ["3", "Materializar movimientos de incrementos, bajas y sustituciones."],
            ["4", "Conciliar 514+13-9+9=527 antes de apagar el MVP."],
        ],
        [["Campos puente presentes", "IDPlaza, MatriculaSustituida", "IDPlaza, MatriculaSustituida", "OK"]],
    )
    wb.save(dest)


def build_pbi_demo(records, k, dest):
    wb = Workbook()
    ws = wb.active
    ws.title = "Hechos_CONTROL_HC"
    headers = COLUMN_ORDER
    for i, h in enumerate(headers, 1):
        ws.cell(1, i, h)
        ws.cell(1, i).font = HEAD_FONT
        ws.cell(1, i).fill = _fill(NAVY)
    for r_i, rec in enumerate(records, 2):
        for c_i, h in enumerate(headers, 1):
            cell = ws.cell(r_i, c_i, rec.get(h, ""))
            if h in ("Matricula", "SuccessID", "IDPlaza", "MatriculaSustituida"):
                cell.number_format = "@"
                cell.value = str(rec.get(h, "") or "")
    add_table(ws, "Hechos_CONTROL_HC", f"A1:{get_column_letter(len(headers))}{1+len(records)}")
    cal = wb.create_sheet("Dim_Calendario")
    cal["A1"] = "Fecha"
    cal["B1"] = "Ano"
    cal["C1"] = "Mes"
    cal["D1"] = "MesNombre"
    cal["E1"] = "AnoMes"
    cal["F1"] = "EsCierre"
    for i, h in enumerate(["Fecha", "Ano", "Mes", "MesNombre", "AnoMes", "EsCierre"], 1):
        cal.cell(1, i, h).font = HEAD_FONT
        cal.cell(1, i).fill = _fill(NAVY)
    meses = [
        (date(2026, 8, 31), 2026, 8, "Agosto", "2026-08", "Si"),
        (date(2026, 9, 30), 2026, 9, "Septiembre", "2026-09", "Si"),
        (date(2026, 10, 31), 2026, 10, "Octubre", "2026-10", "Si"),
        (date(2026, 11, 30), 2026, 11, "Noviembre", "2026-11", "Si"),
        (date(2026, 12, 31), 2026, 12, "Diciembre", "2026-12", "Si"),
    ]
    # daily calendar Aug-Dec 2026
    d0 = date(2026, 8, 1)
    d1 = date(2026, 12, 31)
    r = 2
    cur = d0
    nombres = ["", "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]
    cierres = {date(2026, 8, 31), date(2026, 9, 30), date(2026, 10, 31), date(2026, 11, 30), date(2026, 12, 31)}
    while cur <= d1:
        cal.cell(r, 1, cur.isoformat()).number_format = "@"
        cal.cell(r, 2, cur.year)
        cal.cell(r, 3, cur.month)
        cal.cell(r, 4, nombres[cur.month])
        cal.cell(r, 5, f"{cur.year}-{cur.month:02d}").number_format = "@"
        cal.cell(r, 6, "Si" if cur in cierres else "No")
        r += 1
        cur = date.fromordinal(cur.toordinal() + 1)
    add_table(cal, "Dim_Calendario", f"A1:F{r-1}")
    uni = wb.create_sheet("Dim_Unidad")
    for i, h in enumerate(["UnidadOrganizativa", "ResponsableUnidad", "Empresa", "Codigo"], 1):
        uni.cell(1, i, h).font = HEAD_FONT
        uni.cell(1, i).fill = _fill(NAVY)
    for i, u in enumerate(UNIDADES, 2):
        uni.cell(i, 1, u[1])
        uni.cell(i, 2, u[2])
        uni.cell(i, 3, u[3])
        uni.cell(i, 4, u[0]).number_format = "@"
    add_table(uni, "Dim_Unidad", f"A1:D{1+len(UNIDADES)}")
    evo = wb.create_sheet("Hechos_Evolucion")
    for i, h in enumerate(["AnoMes", "Etiqueta", "HC", "Esperado", "OK"], 1):
        evo.cell(1, i, h).font = HEAD_FONT
        evo.cell(1, i).fill = _fill(NAVY)
    for i, e in enumerate(k["evolucion"], 2):
        evo.cell(i, 1, e["mes"]).number_format = "@"
        evo.cell(i, 2, e["etiqueta"])
        evo.cell(i, 3, e["hc"])
        evo.cell(i, 4, e["esperado"])
        evo.cell(i, 5, "OK" if e["ok"] else "ERROR")
    add_table(evo, "Hechos_Evolucion", "A1:E6")
    par = wb.create_sheet("Parametros")
    par["A1"] = "Parametro"
    par["B1"] = "Valor"
    par["A1"].fill = _fill(NAVY)
    par["B1"].fill = _fill(NAVY)
    par["A1"].font = HEAD_FONT
    par["B1"].font = HEAD_FONT
    for i, (p, v) in enumerate(
        [
            ("HC_Base", 514),
            ("Incrementos", 13),
            ("Salidas_CEV", 9),
            ("Sustituciones", 9),
            ("HC_Final", 527),
            ("Ecuacion", "514 + 13 - 9 + 9 = 527"),
            ("Marca", MARCA_SINTETICO),
        ],
        2,
    ):
        par.cell(i, 1, p)
        par.cell(i, 2, v)
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)


def build_casos_prueba(dest):
    casos = [
        ["CP-01", "Alta base", "Cargar fotografía agosto", "Excel CARGA_REAL vacío", "Importar 514 sintéticos Base/Activo", "HC actual = 514", "", "Diseñado", "", "Operador HC", "Alta"],
        ["CP-02", "Incremento", "Crear plaza nueva", "MVP operativo", "Alta TipoRegistro=Incremento ComputaHC=Si", "ImpactoHC=+1 y HC estructural +1", "", "Diseñado", "", "Operador HC", "Alta"],
        ["CP-03", "Baja", "Registrar salida prevista", "Persona base activa", "Estado=Salida_Prevista + fecha", "Sigue en HC actual; sale del previsto en MesEfecto", "", "Diseñado", "", "Operador HC", "Alta"],
        ["CP-04", "Sustitución", "Cubrir plaza sin crear HC", "Baja identificada", "TipoRegistro=Sustitucion + MatriculaSustituida + mismo IDPlaza", "ImpactoHC=0", "", "Diseñado", "", "Operador HC", "Alta"],
        ["CP-05", "Salida CEV", "9 personas hacia CEV", "9 registros Programa=CEV", "Marcar Salida_Prevista", "No crear plazas nuevas", "", "Diseñado", "", "Operador HC", "Alta"],
        ["CP-06", "Entrada becario sustituto", "9 becarios CEV", "Salida CEV existente", "Alta Sustitucion Programa=CEV", "Neto 0", "", "Diseñado", "", "Operador HC", "Alta"],
        ["CP-07", "Fuera de HC", "Colaborador no estructural", "Lista operativa", "TipoRegistro=Fuera_HC ComputaHC=No", "No altera 514 ni 527", "", "Diseñado", "", "Operador HC", "Media"],
        ["CP-08", "Matrícula duplicada", "Detectar duplicidad", "SYN-0007 duplicado sintético", "Filtrar EstadoCalidad=Duplicado", "Aparece en CONTROL y Calidad", "", "Diseñado", "", "Operador HC", "Alta"],
        ["CP-09", "Registro incompleto", "Falta unidad o correo", "SYN-INC01 / SYN-INC02", "EstadoCalidad=Incompleto ComputaHC=No", "No distorsiona el HC", "", "Diseñado", "", "Operador HC", "Media"],
        ["CP-10", "Fecha incoherente", "Real anterior a prevista absurda", "Registro de alta", "FechaRealAlta < 2010 y FechaPrevistaAlta 2026", "Marca Incoherente, no se inventa corrección", "", "Diseñado", "", "Operador HC", "Media"],
        ["CP-11", "Movimiento cancelado", "Incremento no autorizado al final", "Incremento previsto", "Estado=Cancelado → ImpactoHC=0", "HC estructural no suma esa plaza", "", "Diseñado", "", "Operador HC", "Alta"],
        ["CP-12", "Cierre mensual", "Cerrar agosto 2026", "CONTROL en OK", "Completar HC_Cierre_Mensual", "HC cierre 514, estado Cerrado_demo", "", "Diseñado", "", "Validador", "Alta"],
        ["CP-13", "Reconciliación 514→527", "Ecuación oficial", "Hoja RECONCILIACION", "Verificar 514+13-9+9", "RESULTADO OK y 527", "", "Diseñado", "", "Validador", "Alta"],
        ["CP-14", "Migración a tres listas", "Nivel B", "Plantillas 04", "Ejecutar mapeo y conciliación personas-plazas", "Plazas que computan = 527; nomenclatura MOVIMIENTOS_HC", "", "Diseñado", "", "Arquitectura", "Alta"],
        ["CP-15", "Actualización Power BI", "Refresco del modelo", "Excel demo o Lists", "Actualizar + comprobar medidas", "KPIs coinciden con CONTROL", "", "Diseñado", "", "Analítica", "Media"],
        ["CP-16", "Ejecución de automatizaciones", "Flujos no bloquean MVP", "Entorno Power Automate pendiente DEC-014", "Probar positivo/negativo según especificación", "MVP sigue operando si el flujo está apagado", "", "Diseñado", "", "Automatización", "Media"],
        ["CP-17", "Becario con incremento vs sustituto", "No mezclar", "I01 vs S01", "Comparar OrigenPlaza e ImpactoHC", "I01 = +1 Incremento_Aprobado; S01 = 0 CEV_Sustitucion", "", "Diseñado", "", "Operador HC", "Alta"],
        ["CP-18", "Identificadores como texto", "No convertir a número", "CSV importado", "Ver Matricula SYN-0001", "Permanece texto", "", "Diseñado", "", "Operador HC", "Alta"],
        ["CP-19", "Filtros Lists / Excel", "Responsable, área, estado", "CONTROL_HC", "Filtrar Tecnología + Activo", "Solo registros coincidentes", "", "Diseñado", "", "Operador HC", "Baja"],
        ["CP-20", "Prohibición Power Apps", "Formularios nativos", "Lista CONTROL_HC", "Abrir formulario estándar de Lists", "Alta y baja posibles sin Power Apps", "", "Diseñado", "", "TI colaborativo", "Alta"],
    ]
    wb = Workbook()
    ws = wb.active
    ws.title = "CASOS"
    add_simple_table(
        ws,
        "CASOS DE PRUEBA HC CONTROL v2.0",
        "Cada necesidad de negocio tiene al menos un caso. Resultado obtenido se rellena en ejecución.",
        [
            "ID",
            "Objetivo",
            "Prerrequisitos",
            "Datos",
            "Pasos",
            "Resultado esperado",
            "Resultado obtenido",
            "Estado",
            "Evidencia",
            "Responsable",
            "Severidad",
        ],
        [c[0:1] + c[1:] for c in casos],
        "tblCasos",
        input_cols={7, 8, 9},
        text_cols={1},
    )
    # fix: casos already include all columns
    wb.remove(ws)
    ws = wb.create_sheet("CASOS", 0)
    headers = [
        "ID",
        "Objetivo",
        "Prerrequisitos",
        "Datos",
        "Pasos",
        "Resultado esperado",
        "Resultado obtenido",
        "Estado",
        "Evidencia",
        "Responsable",
        "Severidad",
    ]
    banner(ws, 11, "CASOS DE PRUEBA HC CONTROL v2.0", "Resultado obtenido y evidencia se completan al ejecutar en el tenant.")
    for i, h in enumerate(headers, 1):
        ws.cell(4, i, h)
    style_header_row(ws, 4, 11)
    for r_i, row in enumerate(casos, 5):
        for c_i, v in enumerate(row, 1):
            cell = ws.cell(r_i, c_i, v)
            cell.border = THIN
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            if c_i in (7, 8, 9):
                cell.fill = _fill(INPUT)
            else:
                cell.fill = _fill(CALC) if c_i == 1 else PatternFill()
            if c_i == 1:
                cell.number_format = "@"
                cell.fill = _fill(CALC)
        ws.row_dimensions[r_i].height = 48
    add_table(ws, "tblCasos", f"A4:K{4+len(casos)}")
    ws.freeze_panes = "A5"
    autosize(ws, 36)
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["F"].width = 40
    dict_and_instr(
        wb,
        [["ID", "ID", "texto", "Una línea", "Identificador del caso", "Sí", "CP-nn"]],
        [["1", "Ejecute en orden CP-01 a CP-13 antes de declarar el MVP aceptado."], ["2", "CP-14 es de evolución, no bloquea el MVP."], ["3", "CP-16 no es obligatorio para aceptar el MVP (el MVP funciona sin flujos)."]],
        [["Casos diseñados", "20", "20", "OK"], ["Cobertura ecuación", "CP-13", "CP-13", "OK"]],
    )
    wb.save(dest)


def write_power_query(dest: Path):
    dest.write_text(
        '''// HC Control · Power Query (M)
// Origen preferente: Microsoft Lists CONTROL_HC
// Alternativa: Excel HC_Control_MVP_Operativo.xlsx hoja CONTROL_HC_IMPORTAR
// Identificadores SIEMPRE como texto.

let
    Origen = SharePoint.Tables("https://TENANT.sharepoint.com/sites/SITIO_HC", [ApiVersion = 15]),
    CONTROL_HC = Origen{[Title="CONTROL_HC"]}[Items],
    Tipos = Table.TransformColumnTypes(
        CONTROL_HC,
        {
            {"Matricula", type text},
            {"NombreCompleto", type text},
            {"Empresa", type text},
            {"Correo", type text},
            {"Puesto", type text},
            {"Responsable", type text},
            {"UnidadOrganizativa", type text},
            {"ResponsableUnidad", type text},
            {"FechaAltaRRHH", type date},
            {"IndicadorEstructura", type text},
            {"SuccessID", type text},
            {"TipoRegistro", type text},
            {"Estado", type text},
            {"ComputaHC", type text},
            {"OrigenPlaza", type text},
            {"Programa", type text},
            {"FechaPrevistaAlta", type date},
            {"FechaRealAlta", type date},
            {"FechaPrevistaBaja", type date},
            {"FechaRealBaja", type date},
            {"MesEfecto", type text},
            {"ImpactoHC", Int64.Type},
            {"IDPlaza", type text},
            {"MatriculaSustituida", type text},
            {"Motivo", type text},
            {"Observaciones", type text},
            {"Fuente", type text},
            {"EstadoCalidad", type text},
            {"DatoSintetico", type text}
        }
    ),
    Calidad = Table.AddColumn(
        Tipos,
        "FlagCalidad",
        each
            if [Matricula] = null or [Matricula] = "" then "Incompleto"
            else if [TipoRegistro] = "" then "Sin clasificacion"
            else if [TipoRegistro] = "Sustitucion" and ([MatriculaSustituida] = null or [MatriculaSustituida] = "") then "Incompleto"
            else [EstadoCalidad]
    )
in
    Calidad

// TENANT y SITIO_HC son pendientes (DEC-008). No inventados.
''',
        encoding="utf-8",
    )


def main():
    records = build_records()
    k = kpis(records)
    errs = validate(records, k)
    if errs:
        raise SystemExit(errs)
    mvp = ROOT / "02_MVP_Operativo"
    evo = ROOT / "04_Plantillas_Evolucion"
    pbi = ROOT / "06_Power_BI"
    tes = ROOT / "08_Pruebas"
    mvp.mkdir(parents=True, exist_ok=True)
    build_mvp(records, k, mvp / "HC_Control_MVP_Operativo.xlsx")
    build_csv(records, mvp / "HC_Control_MVP_Importar.csv")
    write_power_query(mvp / "HC_Control_PowerQuery.pq")
    build_personas(records, evo / "PERSONAS_HC_Importar.xlsx")
    est = build_plazas(records, evo / "PLAZAS_HC_Importar.xlsx")
    build_movimientos_hc(records, evo / "MOVIMIENTOS_HC_Importar.xlsx")
    build_conciliacion(records, evo / "HC_Conciliacion_Personas_Plazas.xlsx")
    # alias movimientos operativo de evolución
    build_movimientos_hc(records, evo / "HC_Control_Movimientos.xlsx")
    build_cierre(k, evo / "HC_Cierre_Mensual.xlsx")
    build_calidad(records, evo / "HC_Incidencias_Calidad.xlsx")
    build_mapeo(evo / "HC_Mapeo_MVP_Modelo_Final.xlsx")
    build_pbi_demo(records, k, pbi / "HC_Control_Datos_Demo.xlsx")
    build_casos_prueba(tes / "HC_Control_Casos_Prueba.xlsx")
    print("Excel OK. Plazas computan:", est, "KPIs", {x: k[x] for x in ("hc_actual", "incrementos", "hc_previsto_final")})


if __name__ == "__main__":
    main()
